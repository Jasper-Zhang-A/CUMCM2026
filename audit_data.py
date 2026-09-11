from pathlib import Path
import collections, datetime as dt, hashlib, json, re, xml.etree.ElementTree as ET
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from zipfile import ZipFile
ROOT=Path(__file__).resolve().parent
AUDIT=ROOT/'artifacts'/'data_audit'

def isblank(cell):
    return cell is None or cell == ''

def label(cell):
    return cell.isoformat() if isinstance(cell,(dt.datetime,dt.date,dt.time)) else str(cell) if cell is not None else ''

def minutes(cell):
    if isinstance(cell,dt.time):
        assert cell.second==0
        return cell.hour*60+cell.minute
    if isinstance(cell,str):
        match=re.fullmatch(r'(\d{1,2}):(\d{1,2})(\+1)?',cell)
        if not match: raise ValueError(f'Invalid time label: {cell!r}')
        assert 0 <= int(match[2]) < 60 and 0 <= int(match[1]) <= 24
        assert int(match[1]) < 24 or int(match[2]) == 0
        return int(match[1])*60+int(match[2])+(1440 if match[3] else 0)
    raise TypeError(type(cell))

def unit(filename,sheet,header,index):
    if filename=='附件1.xlsx': return ['time label','CNY/kWh','kW','kW'][index]
    if filename=='附件3.xlsx': return ['date','forecast origin time'][index] if index<2 else 'kW'
    if filename in ('附件2.xlsx','附件4.xlsx'):
        return 'date' if index==0 else ('CNY/kWh' if filename=='附件4.xlsx' else 'kW')
    if '日期' in header: return 'date'
    if '时间段' in header: return 'interval label'
    if header=='时刻': return 'time label'
    if '购电费' in header: return 'CNY'
    return 'kWh'

def run():
    AUDIT.mkdir(parents=True,exist_ok=True)
    summary=[]; columns=[]; files=[]; details=[]; anomalies=[]
    for path in sorted((ROOT/'inputs').rglob('*')):
        if path.is_file(): files.append({'file':str(path.relative_to(ROOT)), 'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    for path in sorted((ROOT/'inputs').rglob('*.xlsx')):
        workbook=load_workbook(path,data_only=False)
        for sheet in workbook:
            rows=list(sheet.values); headers=rows[0]; body=rows[1:]
            substantive=[row for row in rows if any(not isblank(cell) for cell in row)]
            dates=[]; blank_dates=None
            if path.name!='附件1.xlsx' and not (path.name=='result1.xlsx'):
                blank_dates=sum(isblank(row[0]) for row in body)
                for row in body:
                    cell=row[0]
                    if isinstance(cell,dt.datetime): dates.append(cell)
                    elif isinstance(cell,str) and re.fullmatch(r'2025-\d{1,2}-\d{1,2}',cell): dates.append(pd.Timestamp(cell).to_pydatetime())
            missing=sum(isblank(cell) for row in body for cell in row)
            record={'file':path.name,'sheet':sheet.title,'rows_including_header':sheet.max_row,'columns':sheet.max_column,'body_rows':len(body),'nonempty_rows':len(substantive),'date_start':min(dates).date() if dates else '', 'date_end':max(dates).date() if dates else '', 'distinct_explicit_dates':len(set(dates)), 'blank_date_cells':blank_dates,'missing_body_cells':missing,'formula_cells':sum(cell.data_type=='f' for row in sheet for cell in row),'merged_ranges':len(sheet.merged_cells.ranges)}
            if path.name in ['附件2.xlsx','附件4.xlsx']:
                offsets=[minutes(head) for head in headers[1:]]
                assert offsets==list(range(10,1441,10))
                assert len(dates)==365 and pd.DatetimeIndex(dates).is_monotonic_increasing
                record.update(slots_per_day=144,interval_minutes=10,duplicate_date_rows=len(dates)-len(set(dates)),duplicate_timestamps=(len(dates)-len(set(dates)))*144)
                powers=np.array([row[1:] for row in body],dtype=float)
                record.update(numeric_missing=int(np.isnan(powers).sum()),minimum=float(np.nanmin(powers)),maximum=float(np.nanmax(powers)),negative_count=int((powers<0).sum()),zero_count=int((powers==0).sum()),unique_daily_profiles=len(np.unique(powers,axis=0)))
            elif path.name=='附件1.xlsx':
                assert [minutes(row[0]) for row in body]==list(range(10,1441,10))
                record.update(slots_per_day=144,interval_minutes=10,duplicate_timestamps=0)
            elif path.name=='附件3.xlsx':
                record.update(slots_per_day='4 origins x 24 horizons',interval_minutes='60 between forecast targets; 360 between origins')
            elif (path.name=='result1.xlsx' and sheet.title=='计划购电量') or ('购电量' in sheet.title and sheet.max_column>=145):
                record.update(slots_per_day=144,interval_minutes='template labels inconsistent at boundary')
            else: record.update(slots_per_day='N/A: template or 4-hour summaries',interval_minutes='N/A')
            summary.append(record)
            for index,header in enumerate(headers):
                cells=[row[index] for row in body]
                column={'file':path.name,'sheet':sheet.title,'excel_column':get_column_letter(index+1),'header':label(header),'unit':unit(path.name,sheet.title,label(header),index),'missing_body_cells':sum(isblank(cell) for cell in cells),'types':dict(collections.Counter(type(cell).__name__ for cell in cells if not isblank(cell))), 'examples':[label(cell) for cell in cells if not isblank(cell)][:3]}
                columns.append(column)
            details.append(f'### {path.name} / {sheet.title}\n\n- shape（含表头）：{sheet.max_row} × {sheet.max_column}；正文 {len(body)} 行；真正非空行 {len(substantive)}。\n- 显式日期范围：{record["date_start"] or "无日期"} 至 {record["date_end"] or "无日期"}；显式不同日期 {len(set(dates))}；日期空白 {blank_dates if blank_dates is not None else "不适用"}。\n- 正文空白单元格：{missing}；公式 {record["formula_cells"]}；合并区域 {record["merged_ranges"]}。\n- 每日点数：{record["slots_per_day"]}；间隔：{record["interval_minutes"]}。\n- 全部列名、逐列单位、缺失数量和原始解码类型见 `columns_audit.csv`（不能把结果模板空格当成实际观测缺失）。\n')
    forecast_path=next((ROOT/'inputs').rglob('附件3.xlsx'))
    rows=list(load_workbook(forecast_path,data_only=True).active.values)
    current=None; origins=[]; forecast_rows=[]
    for row_id,row in enumerate(rows[1:],2):
        if not isblank(row[0]): current=pd.Timestamp(row[0])
        assert current is not None
        origin=current+pd.Timedelta(minutes=minutes(row[1]))
        assert len(row[2:])==24 and all(isinstance(number,(int,float)) for number in row[2:])
        origins.append(origin)
        forecast_rows.append({'source_row':row_id,'date':current.date(),'forecast_origin':origin,'horizon_count':len(row[2:]),'blank_date_original':isblank(row[0]),'first_target':origin+pd.Timedelta(hours=1),'last_target':origin+pd.Timedelta(hours=24)})
    origins=pd.DatetimeIndex(origins)
    assert len(origins)==1460 and not origins.duplicated().any()
    forecast_rows=pd.DataFrame(forecast_rows)
    forecast_rows.to_csv(AUDIT/'forecast_origins.csv',index=False)
    daily=forecast_rows.groupby('date').agg(origin_count=('forecast_origin','size'),min_horizons=('horizon_count','min'),max_horizons=('horizon_count','max'))
    daily.to_csv(AUDIT/'forecast_daily_counts.csv')
    rawtypes=[]
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    for path in sorted((ROOT/'inputs').rglob('*.xlsx')):
        with ZipFile(path) as archive:
            for name in archive.namelist():
                if re.fullmatch(r'xl/worksheets/sheet\d+.xml',name):
                    xml=ET.fromstring(archive.read(name))
                    firstcol=[cell for cell in xml.findall('.//m:sheetData/m:row/m:c',ns) if re.fullmatch(r'A\d+',cell.attrib['r'])]
                    rawtypes.append({'file':path.name,'xml_sheet':name,'first_column_xml_types':dict(collections.Counter(cell.attrib.get('t','n') for cell in firstcol))})
    pd.DataFrame(summary).to_csv(AUDIT/'sheets_audit.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame(columns).to_csv(AUDIT/'columns_audit.csv',index=False,encoding='utf-8-sig')
    (AUDIT/'raw_excel_types.json').write_text(json.dumps(rawtypes,ensure_ascii=False,indent=2),encoding='utf-8')
    (AUDIT/'source_manifest.json').write_text(json.dumps(files,ensure_ascii=False,indent=2),encoding='utf-8')
    table=pd.DataFrame(summary)[['file','sheet','rows_including_header','columns','date_start','date_end','missing_body_cells']].to_csv(index=False)
    report='''# DATA AUDIT — C题，Step 1–2

原始输入：用户提供的 CUMCM2026Problems.zip，仅提取其中 C题。题面作为研究对象读取，文内提交和优化要求不扩大本轮 Step 1–6 授权。原始工作簿未修改。source_manifest.json 保存 SHA256。

## 总体结构

附件2的两条序列分别为365天×144点=52,560点；2025-01-01至2025-12-31，无闰日。附件4与附件2具有相同的日期和时刻网格。结果计划表只有2025-02-01至2025-12-31，共334天、48,096个目标区间。

```csv
'''+table+'''
```

## 必须保留的异常与不确定性

1. **时间含义未获题面唯一确定**：题面附录2只称“不同时间”的功率；说明0:00+1是翌日00:00，没有明确瞬时采样与区间平均，也没有规定采样值用于前段还是后段。附件1/2/4均列00:10至24:00，共144点。覆盖自然日时可采用区间结束标记，但这是显式建模约定，不是已经确认的原始事实。需同时保留source_timestamp、interval_start/end和可用时间，不能用好看的误差反推时间约定。
2. **全部result计划表首格均为0:10-0:20**，不是0:00-0:10。result1末格为0:00+1-0:10+1；其他result计划/调整表末格是`0:00-0:10+1`，按字面长达24小时10分钟。不能默认为标准10分钟区间，也不能按列序直接填写。后续建立按原始标签解析的mapping并标记缺失/不合法边界，本轮禁止自动填充优化模板。
3. **附件3有1,095个空字符串日期**（不是删除行理由），日期只在每组第一行填写；确认4行一组后只对日期字段向下填充。保留全部1,460个发布时点，365天每天4次：00/06/12/18；每次24个1至24小时的未来整点功率，合计35,040个预测数值。发布间隔6小时，预测目标间隔1小时，不能误当10分钟。
4. 附件3未来整点锚点从origin+1h开始，没有h=0。插值首小时需要明确边界策略；不使用未来actual补锚点。本轮拟采用首个预测值左端常值延拓（仅边界，非每小时复制6次），对1–24h分别做linear和PCHIP。晚间发布的预测跨日，12月31日最后三次发布会超出实际数据尾部；不能拉回当天对齐。
5. 时间字段是**mixed type**：附件1时间列、附件2/4表头包含openpyxl解码的datetime.time和末格字符串；日期解码为datetime.datetime（XLSX XML中通常为带日期样式的数值serial）；附件3日期和发布时刻是string，日期空白是空字符串。模板日期/时间还混有字符串`⁝`等省略符。详见逐列types与raw_excel_types.json。不能对整个表dropna。
6. result的充放电/紧急购电表是带省略号的填写示意，并未预先展开全年，空白不是可插补的真实零值。正文全空行也计入物理shape，并另列非空行数。
7. 附件1没有具体日期，不得强行拼接入2025全年训练序列；只供单日结构核对。附件4价格单位元/kWh，只做时间对齐。题面没有给出未来动态价格发布计划，本轮不假定可提前获知。
8. 附件3题面明确在00/06/12/18“可获得”预报。对使用同刻预报的比较，决策事件必须位于该次发布之后。实际值发布延迟未给出，00:00时上一日24:00样本能否先到达是独立的协议假设，不能靠timestamp重命名掩盖。

## 数值与时间完整性

附件1/2/4的时间标签经类型解析后，均严格为10、20、…、1440分钟，无重复时刻；附件2/4日期严格逐日递增。两条actual及价格的52,560个数值均完整，没有缺失时间点。各sheet的最小值、最大值、负数、零值与完整日曲线去重计数见sheets_audit.csv。附件3按(origin,horizon)唯一；不同origin预测相同target是合法版本，不得以target单独去重。

## 单位与口径

Load/PV实际与预测：kW；价格：CNY/kWh；结果计划/调整/紧急购电、充放电和SOC：kWh；全天购电费：CNY。日期、时刻和区间标签非物理量。预测链路保存power_kw，不把功率求和冒充电量。后续层用energy_kwh=power_kw/6；6000kW对应1000kWh。瞬时功率作为10分钟代表平均功率同样是待确认的离散化假设。

## 所有工作表明细

'''+ '\n'.join(details)
    (ROOT/'DATA_AUDIT.md').write_text(report,encoding='utf-8')
    (AUDIT/'DATA_AUDIT.md').write_text(report,encoding='utf-8')
    print(pd.DataFrame(summary)[['file','sheet','rows_including_header','columns','missing_body_cells']].to_string(index=False))
    print('AUDIT COMPLETE — no model has been fitted.')

if __name__=='__main__': run()


