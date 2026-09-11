from pathlib import Path
import collections, datetime as dt, hashlib, importlib.metadata, json, platform
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from audit_data import isblank, label, minutes, ROOT, AUDIT

def markdown_table(frame):
    frame=frame.fillna('')
    return '| '+' | '.join(frame.columns)+' |\n|'+'|'.join(['---']*len(frame.columns))+'|\n'+'\n'.join('| '+' | '.join(str(cell).replace('|','/') for cell in row)+' |' for row in frame.itertuples(index=False,name=None))

def run():
    time_cells=[]; label_checks=[]; daily_counts=[]; template_intervals=[]; diagnostics=[]; checks=[]
    numeric_stats=[]; date_fields=[]; type_summary=[]
    def check(name,condition,detail):
        checks.append({'check':name,'passed':bool(condition),'detail':detail})
        assert condition, f'{name}: {detail}'
    for path in sorted((ROOT/'inputs').rglob('*.xlsx')):
        workbook=load_workbook(path,data_only=True)
        for sheet in workbook:
            rows=list(sheet.values); header=rows[0]
            is_wide=path.name in ['附件2.xlsx','附件4.xlsx']
            if path.name=='附件1.xlsx' or is_wide:
                cells=list(sheet['A'])[1:] if path.name=='附件1.xlsx' else list(sheet[1])[1:]
                for cell in cells:
                    time_cells.append({'file':path.name,'sheet':sheet.title,'cell':cell.coordinate,'raw_label':label(cell.value),'decoded_type':type(cell.value).__name__,'number_format':cell.number_format,'minute_offset':minutes(cell.value),'semantic_status':'unresolved point/interval meaning'})
                offsets=[minutes(cell.value) for cell in cells]
                check(f'{path.name}/{sheet.title}/10min_grid',offsets==list(range(10,1441,10)),'144 distinct offsets 10..1440')
                label_checks.append({'file':path.name,'sheet':sheet.title,'decoded_time_types':dict(collections.Counter(type(cell.value).__name__ for cell in cells)),'point_count':len(offsets),'duplicate_offsets':len(offsets)-len(set(offsets)),'first_source_offset_minutes':min(offsets),'last_source_offset_minutes':max(offsets),'all_gaps_10min':bool(np.all(np.diff(offsets)==10))})
            if is_wide:
                dates=pd.DatetimeIndex([row[0] for row in rows[1:]])
                check(f'{path.name}/{sheet.title}/date_grid',dates.equals(pd.date_range('2025-01-01','2025-12-31')), '365 exact consecutive natural-day row labels')
                stamps=[]
                for row in rows[1:]:
                    day=pd.Timestamp(row[0]); powers=np.asarray(row[1:],dtype=float)
                    stamps.extend(day+pd.to_timedelta(offsets,unit='min'))
                    daily_counts.append({'file':path.name,'sheet':sheet.title,'date':day.date(),'expected_points':144,'numeric_points':int(np.isfinite(powers).sum()),'missing_numeric':int(np.isnan(powers).sum()),'duplicate_offset_count':len(offsets)-len(set(offsets))})
                stamps=pd.DatetimeIndex(stamps)
                check(f'{path.name}/{sheet.title}/source_timestamp_unique',not stamps.duplicated().any(),'52,560 source timestamps, keyed by series')
                check(f'{path.name}/{sheet.title}/continuous_grid',bool((stamps[1:] - stamps[:-1] == pd.Timedelta(minutes=10)).all()),'source timestamps from 2025-01-01 00:10 to 2026-01-01 00:00')
            if path.name.startswith('result') and (sheet.title in ['计划购电量','调整购电量']):
                cells=list(sheet['A'])[1:] if path.name=='result1.xlsx' else list(sheet[1])[1:145]
                offsets=[]
                for position,cell in enumerate(cells):
                    start,end=cell.value.split('-')
                    start_min=minutes(start); end_min=minutes(end)
                    duration=end_min-start_min
                    status='valid_10min_in_natural_day' if duration==10 and 0<=start_min<1440 and end_min<=1440 else ('outside_natural_day' if duration==10 else 'invalid_duration')
                    template_intervals.append({'file':path.name,'sheet':sheet.title,'cell':cell.coordinate,'source_order_only':position,'raw_interval_label':cell.value,'literal_start_minute':start_min,'literal_end_minute':end_min,'literal_duration_minute':duration,'status':status})
                    # Diagnostic only: row ordering is not an adopted mapping.
                    diagnostics.append({'file':path.name,'sheet':sheet.title,'source_column':get_column_letter(position+2) if path.name!='result1.xlsx' else f'A{position+2}', 'source_label_offset_minute':(position+1)*10,'slot_id':'','result_column':cell.coordinate,'result_label':cell.value,'start_minus_same_order_source_label_minute':start_min-(position+1)*10,'mapping_status':'UNRESOLVED_DO_NOT_EXPORT'})
                    offsets.append(start_min)
                label_checks.append({'file':path.name,'sheet':sheet.title,'decoded_time_types':{'str':len(cells)},'point_count':len(cells),'duplicate_offsets':len(offsets)-len(set(offsets)),'first_source_offset_minutes':min(offsets),'last_source_offset_minutes':max(offsets),'all_gaps_10min':bool(np.all(np.diff(offsets)==10))})
            # Fully audit column A/B time/date fields, including blank and placeholder encodings.
            if path.name!='附件1.xlsx':
                time_indices=[0,1] if path.name=='附件3.xlsx' else [index for index,head in enumerate(header) if isinstance(head,str) and any(part in head for part in ['日期','时间段','时刻','预报时刻'])]
                for index in time_indices:
                    field_cells=[row[index] for row in rows[1:]]
                    nonblank=[cell for cell in field_cells if not isblank(cell)]
                    type_summary.append({'file':path.name,'sheet':sheet.title,'column':get_column_letter(index+1),'header':label(header[index]),'empty_none':sum(cell is None for cell in field_cells),'empty_string':sum(isinstance(cell,str) and cell=='' for cell in field_cells),'types':dict(collections.Counter(type(cell).__name__ for cell in nonblank)),'placeholder_count':sum(cell=='⁝' for cell in nonblank)})
            if path.name in ['附件1.xlsx','附件3.xlsx']:
                names=header[1:] if path.name=='附件1.xlsx' else ['PV forecasts (24 horizons combined)']
                blocks=[(name,[row[index] for row in rows[1:]]) for index,name in enumerate(names,1)] if path.name=='附件1.xlsx' else [(names[0],[cell for row in rows[1:] for cell in row[2:]])]
                for name,powers in blocks:
                    powers=np.asarray(powers,dtype=float)
                    numeric_stats.append({'file':path.name,'series':name,'count':len(powers),'missing_count':int(np.isnan(powers).sum()),'nonfinite_count':int((~np.isfinite(powers)).sum()),'min':powers.min(),'max':powers.max(),'negative_count':int((powers<0).sum()),'zero_count':int((powers==0).sum())})
                    check(f'{path.name}/{name}/finite',bool(np.isfinite(powers).all()),f'{len(powers)} numeric observations')
    frame=pd.DataFrame(template_intervals)
    summary=frame.groupby(['file','sheet','status']).size().reset_index(name='label_count')
    path=next((ROOT/'inputs').rglob('附件3.xlsx'))
    sheet=load_workbook(path,data_only=True).active
    rows=list(sheet.values); current=None; observed_date_groups=[]; horizons=[]
    for source_row,row in enumerate(rows[1:],2):
        if not isblank(row[0]):
            current=pd.Timestamp(row[0]); observed_date_groups.append(source_row)
        origin=current+pd.Timedelta(minutes=minutes(row[1]))
        for hour,power in enumerate(row[2:],1):
            horizons.append({'source_row':source_row,'source_column':get_column_letter(hour+2),'forecast_origin':origin,'horizon_hour':hour,'target_source_timestamp':origin+pd.Timedelta(hours=hour),'forecast_power_kw':power})
    check('附件3/date_blocks',observed_date_groups==list(range(2,1462,4)),'nonblank dates exactly on the first row of every 4-row group; only date ffill permitted')
    for day_index in range(365):
        check(f'附件3/day_{day_index+1}/release_hours',[minutes(row[1]) for row in rows[1+day_index*4:5+day_index*4]]==[0,360,720,1080],'four ordered release hours')
    check('附件3/horizon_labels',list(rows[0][2:])==[f'预报{hour}小时' for hour in range(1,25)],'24 exact ordinal hour labels')
    horizons=pd.DataFrame(horizons)
    duplicate_keys=int(horizons.duplicated(['forecast_origin','horizon_hour']).sum())
    late=int((horizons.target_source_timestamp>pd.Timestamp('2026-01-01')).sum())
    check('附件3/origin_horizon_unique',duplicate_keys==0,'35,040 unique forecast version keys; target duplicates are legitimate revisions')
    check('附件3/future_target',bool((horizons.target_source_timestamp>horizons.forecast_origin).all()),'targets strictly after each release')
    for filename,frame_out in [('time_field_cells.csv',pd.DataFrame(time_cells)),('time_label_summary.csv',pd.DataFrame(label_checks)),('daily_point_counts.csv',pd.DataFrame(daily_counts)),('template_intervals.csv',pd.DataFrame(template_intervals)),('alignment_diagnostic_NOT_A_MAPPING.csv',pd.DataFrame(diagnostics)),('template_interval_summary.csv',summary),('time_date_field_types.csv',pd.DataFrame(type_summary)),('additional_numeric_stats.csv',pd.DataFrame(numeric_stats)),('attachment3_hourly_audit.csv',horizons)]:
        frame_out.to_csv(AUDIT/filename,index=False,encoding='utf-8-sig')
    # Verify preserved input hashes against extracted original bytes in ZIP, including PDF.
    from zipfile import ZipFile
    cfg=json.loads((ROOT/'configs/audit.json').read_text(encoding='utf-8-sig'))
    with ZipFile(cfg['source_zip']) as archive:
        for path in (ROOT/'inputs').rglob('*'):
            if path.is_file():
                member=path.relative_to(ROOT/'inputs').as_posix()
                check('preserved/'+member,hashlib.sha256(path.read_bytes()).digest()==hashlib.sha256(archive.read(member)).digest(),'original file bytes unchanged')
    report=(ROOT/'DATA_AUDIT.md').read_text(encoding='utf-8')
    report=report.replace('原始输入：','**执行范围：按用户后续明确选择，仅完成审计。canonical interval语义未冻结；不训练、不生成预测、不运行forecasting leakage测试。**\n\n原始输入：',1)
    report=report.replace('2025-02-01至2025-12-31，共334天、48,096个目标区间。','2025-02-01至2025-12-31，共334天、48,096个待预测位置；区间语义未确定。')
    report=report.replace('本轮拟采用首个预测值左端常值延拓（仅边界，非每小时复制6次），对1–24h分别做linear和PCHIP。','后续可比较首个预测值左端常值延拓、可用历史锚点等边界方案；本轮按用户要求仅指出问题，不插值、不确定默认方案。')
    report=report.replace('附件1时间列、附件2/4表头包含openpyxl解码的datetime.time和末格字符串；','附件1时间列前60个标签（00:10–10:00）为openpyxl解码的datetime.time，后84个（10:10–24:00）均为string；附件2/4表头前143个为datetime.time、末格为string；')
    table_start=report.index('```csv\n'); table_end=report.index('\n```',table_start)+4
    sheets=pd.read_csv(AUDIT/'sheets_audit.csv').fillna('')
    overview=sheets[['file','sheet','rows_including_header','columns','date_start','date_end','missing_body_cells']].copy()
    overview.columns=['文件','工作表','行(含表头)','列','开始日期','结束日期','正文空白']
    report=report[:table_start]+markdown_table(overview)+report[table_end:]
    appendix=f'''\n\n## 标签逐格复核及补充结果

{markdown_table(summary.rename(columns={'file':'文件','sheet':'工作表','status':'字面状态','label_count':'标签数'}))}

状态含义：valid_10min_in_natural_day为自然日内合法10分钟标签；outside_natural_day为区间位于自然日以外；invalid_duration为标签字面时长不等于10分钟。每张计划/调整表都有143个自然日内合法区间，但均缺少[00:00,00:10)。result1另有1个次日区间；其余6张表各有1个末格时长异常。

`template_intervals.csv`保留每个Excel地址、原始标签、字面起止分钟及异常状态；`alignment_diagnostic_NOT_A_MAPPING.csv`是同序位置差异的诊断，不是可执行mapping，slot_id有意留空。只有时间解释获确认后才建立正式source_column/slot_id/result_column映射。

- 模板另外有6处单数字分钟标签：result2/3/4-2/4-3的计划表及result3/4-3的调整表AQ1均为`7:0-7:10`。解析器显式支持一位或两位分钟，并校验分钟0–59；原字符串保留，不修改工作簿。
- 时间字段没有时区信息。本轮保留原始本地墙钟时间，不擅自赋予时区或夏令时规则；后续外部数据接入需要统一时区。
- 附件2/4原始源时间戳范围：2025-01-01 00:10至2026-01-01 00:00，按源日期字段分组时每天144点。若直接按timestamp.date分组，首日143点、尾日1点；这不意味着原始数据缺日。date必须保留来源日期，不能误把末点丢到训练的次日。
- 附件3最早目标为2025-01-01 01:00，最晚目标为2026-01-01 18:00。有{late}个预测锚点晚于附件2最后一个源时刻，缺少对应ground truth；不能填0或强行回卷。它们来自12月31日06/12/18时发布的预测，分别6/12/18点。
- 附件3数值无缺失，(origin,horizon)重复数{duplicate_keys}；相同target的多个版本是预期行为。每一天的发布次数和每次horizon数均逐行核对，而非只检查首日。
- 附件1同一时间列在10:10发生time→string类型切换。仅支持datetime.time或只特判最后一个字符串的读取器将失败，详见time_field_cells.csv。
- 负载范围1995.7176–7978.8849 kW，PV范围0–10216.2 kW，价格0.0076–1.7936元/kWh。PV有23540个零值（{23540/52560:.2%}）。这只是审计统计，未来不能用于fit/scaler、日照阈值或模型选择。
- Load/PV各有365条不同的完整日曲线。52,560个点不等于52,560个独立样本，也不能把365天视为统计独立。全年只有52个完整7天周期加1天；2月1日前只有31个日周期，实际可见的最后一格还取决于00:00接收协议。

### 额外数值审计

{markdown_table(pd.DataFrame(numeric_stats))}

## 检查状态与复现

本轮运行的是原始数据结构和边界核验，不是预测泄漏测试。`structural_checks.json`保存{len(checks)}项断言结果，全部通过；通过的是与本报告中异常相容的读取完整性检查，**不表示模板标签正确**。

- DATA AUDIT：完成。
- Canonical interval定义 / 正式mapping：等待澄清。
- tests/test_no_future_leakage.py：尚未实现和执行，不能报告通过。
- B0–B4、metrics、预测图：未运行。
- DLinear / PatchTST：未启动。

复现命令：`.\\.venv\\Scripts\\python.exe run_audit.py`。配置在configs/audit.json，依赖与运行环境在requirements-audit.txt及artifacts/data_audit/environment.json。运行入口只审计，无法启动模型训练。原始PDF和全部9个Excel文件已逐字节核对与ZIP一致。
'''
    report+=appendix
    (ROOT/'DATA_AUDIT.md').write_text(report,encoding='utf-8')
    (AUDIT/'DATA_AUDIT.md').write_text(report,encoding='utf-8')
    (AUDIT/'structural_checks.json').write_text(json.dumps({'scope':'DATA AUDIT ONLY - NOT FORECAST LEAKAGE TESTS','passed':len(checks),'failed':0,'checks':checks},ensure_ascii=False,indent=2),encoding='utf-8')
    environment={'python':platform.python_version(),'platform':platform.platform(),'packages':{name:importlib.metadata.version(name) for name in ['numpy','pandas','openpyxl','pypdf','pymupdf']},'pytorch':'not used','TSLib_commit':'not installed / not run','CUDA_GPU':'not used / not probed','random_seed':cfg['random_seed'],'training_time_seconds':0,'parameter_count':0,'stage':'audit_only'}
    (AUDIT/'environment.json').write_text(json.dumps(environment,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'EXTRA AUDIT COMPLETE: {len(checks)} structural checks passed; {late} forecast anchors beyond actual coverage; no models run.')

if __name__=='__main__': run()



