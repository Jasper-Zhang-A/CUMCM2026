from pathlib import Path
import zipfile, json
from pypdf import PdfReader
from openpyxl import load_workbook
root=Path.cwd()
with zipfile.ZipFile(r'C:\Users\lenovo\Desktop\CUMCM2026Problems.zip') as z:
    for info in z.infolist():
        if info.filename.startswith('C题/') and not info.is_dir():
            dest=root/'inputs'/info.filename
            dest.parent.mkdir(parents=True,exist_ok=True)
            dest.write_bytes(z.read(info))
pdf=next((root/'inputs').rglob('*.pdf'))
pages=[p.extract_text() for p in PdfReader(pdf).pages]
(root/'problem_text.txt').write_text('\n\n'.join(f'PAGE {i+1}\n{t}' for i,t in enumerate(pages)),encoding='utf-8')
print((root/'problem_text.txt').read_text(encoding='utf-8'))
for path in sorted((root/'inputs').rglob('*.xlsx')):
    print('\nFILE',path.name)
    wb=load_workbook(path,read_only=False,data_only=False)
    for ws in wb:
        print('SHEET',ws.title,'SHAPE',ws.max_row,ws.max_column,'MERGES',list(ws.merged_cells.ranges)[:8])
        for row in list(ws.values)[:5]:
            print(repr(row[:8]),'... LAST',repr(row[-3:]))
        print('LAST_ROW',repr(list(ws.values)[-1][:5]))
